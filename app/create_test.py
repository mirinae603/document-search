#!/usr/bin/env python3
"""
Generate sample test files for document extraction testing
"""
import os
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import io


def create_test_files_directory():
    """Create test files directory structure"""
    test_dir = Path("./test_files")
    test_dir.mkdir(exist_ok=True)
    return test_dir


def create_sample_txt(test_dir: Path):
    """Create sample text file"""
    content = """This is a sample text document for extraction testing.

It contains multiple paragraphs with various formatting.
Special characters: @#$%^&*()
Numbers: 1234567890

This file tests basic text extraction capabilities."""
    
    with open(test_dir / "sample.txt", "w") as f:
        f.write(content)


def create_sample_csv(test_dir: Path):
    """Create sample CSV file"""
    content = """Name,Age,City,Department
John Doe,30,New York,Engineering
Jane Smith,25,San Francisco,Marketing
Bob Johnson,35,Chicago,Sales
Alice Williams,28,Boston,HR"""
    
    with open(test_dir / "sample.csv", "w") as f:
        f.write(content)


def create_sample_html(test_dir: Path):
    """Create sample HTML file"""
    content = """<!DOCTYPE html>
<html>
<head>
    <title>Sample HTML Document</title>
</head>
<body>
    <h1>Document Extraction Test</h1>
    <p>This is a <strong>sample HTML</strong> document for testing.</p>
    <ul>
        <li>Item 1</li>
        <li>Item 2</li>
        <li>Item 3</li>
    </ul>
    <p>It contains various HTML elements including <em>emphasis</em> and <code>code</code>.</p>
</body>
</html>"""
    
    with open(test_dir / "sample.html", "w") as f:
        f.write(content)


def create_sample_image_with_text(test_dir: Path):
    """Create sample images with text for OCR testing"""
    
    # Create a simple image with text
    img = Image.new('RGB', (800, 400), color='white')
    draw = ImageDraw.Draw(img)
    
    # Try to use a font, fall back to default if not available
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 40)
        font_small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24)
    except:
        font = ImageFont.load_default()
        font_small = ImageFont.load_default()
    
    # Draw text
    draw.text((50, 50), "OCR Test Document", fill='black', font=font)
    draw.text((50, 120), "This image contains text that should be", fill='black', font=font_small)
    draw.text((50, 160), "extracted using Optical Character Recognition.", fill='black', font=font_small)
    draw.text((50, 220), "Numbers: 1234567890", fill='black', font=font_small)
    draw.text((50, 260), "Special chars: @#$%^&*()", fill='black', font=font_small)
    
    # Save in different formats
    img.save(test_dir / "sample.png", "PNG")
    img.save(test_dir / "sample.jpg", "JPEG")
    img.save(test_dir / "sample.bmp", "BMP")
    
    print("✅ Created sample images (PNG, JPG, BMP)")


def create_sample_docx(test_dir: Path):
    """Create sample DOCX file using python-docx"""
    try:
        from docx import Document
        
        doc = Document()
        doc.add_heading('Document Extraction Test', 0)
        
        doc.add_paragraph('This is a sample Microsoft Word document for testing extraction capabilities.')
        
        doc.add_heading('Section 1: Introduction', level=1)
        doc.add_paragraph('This document contains various formatting elements to test extraction quality.')
        
        doc.add_heading('Section 2: Features', level=1)
        doc.add_paragraph('The document includes:')
        
        # Add bullet points
        doc.add_paragraph('Headings and paragraphs', style='List Bullet')
        doc.add_paragraph('Bold and italic text', style='List Bullet')
        doc.add_paragraph('Tables and lists', style='List Bullet')
        
        # Add a table
        table = doc.add_table(rows=3, cols=3)
        table.style = 'Light Grid Accent 1'
        
        # Header row
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Name'
        hdr_cells[1].text = 'Age'
        hdr_cells[2].text = 'City'
        
        # Data rows
        row1_cells = table.rows[1].cells
        row1_cells[0].text = 'John'
        row1_cells[1].text = '30'
        row1_cells[2].text = 'NYC'
        
        row2_cells = table.rows[2].cells
        row2_cells[0].text = 'Jane'
        row2_cells[1].text = '25'
        row2_cells[2].text = 'SF'
        
        doc.add_paragraph()
        doc.add_paragraph('This document helps verify that the extraction system can handle complex formatting.')
        
        doc.save(test_dir / "sample.docx")
        print("✅ Created sample DOCX")
        
    except ImportError:
        print("⚠️  python-docx not installed. Skipping DOCX creation.")
        print("   Install with: pip install python-docx")


def create_sample_xlsx(test_dir: Path):
    """Create sample XLSX file using openpyxl"""
    try:
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Sample Data"
        
        # Add headers
        headers = ['Employee ID', 'Name', 'Department', 'Salary', 'Start Date']
        ws.append(headers)
        
        # Add sample data
        data = [
            [1001, 'John Doe', 'Engineering', 75000, '2020-01-15'],
            [1002, 'Jane Smith', 'Marketing', 65000, '2019-03-20'],
            [1003, 'Bob Johnson', 'Sales', 70000, '2021-06-10'],
            [1004, 'Alice Williams', 'HR', 60000, '2018-11-05'],
            [1005, 'Charlie Brown', 'Engineering', 80000, '2020-08-22'],
        ]
        
        for row in data:
            ws.append(row)
        
        # Add a second sheet
        ws2 = wb.create_sheet("Summary")
        ws2.append(['Total Employees', 5])
        ws2.append(['Average Salary', 70000])
        ws2.append(['Departments', 4])
        
        wb.save(test_dir / "sample.xlsx")
        print("✅ Created sample XLSX")
        
    except ImportError:
        print("⚠️  openpyxl not installed. Skipping XLSX creation.")
        print("   Install with: pip install openpyxl")


def create_sample_pdf(test_dir: Path):
    """Create sample PDF using reportlab"""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        
        # Create PDF
        pdf_file = test_dir / "sample.pdf"
        doc = SimpleDocTemplate(str(pdf_file), pagesize=letter)
        
        # Container for the 'Flowable' objects
        elements = []
        
        # Define styles
        styles = getSampleStyleSheet()
        title_style = styles['Title']
        heading_style = styles['Heading1']
        normal_style = styles['Normal']
        
        # Add content
        elements.append(Paragraph("Document Extraction Test PDF", title_style))
        elements.append(Spacer(1, 0.2*inch))
        
        elements.append(Paragraph("Introduction", heading_style))
        elements.append(Paragraph(
            "This is a sample PDF document created for testing document extraction capabilities. "
            "It contains various elements including text, tables, and formatting.",
            normal_style
        ))
        elements.append(Spacer(1, 0.2*inch))
        
        elements.append(Paragraph("Features", heading_style))
        elements.append(Paragraph(
            "This document demonstrates:<br/>"
            "• Multiple paragraphs with formatting<br/>"
            "• Tables with structured data<br/>"
            "• Headers and sections<br/>"
            "• Various text styles",
            normal_style
        ))
        elements.append(Spacer(1, 0.2*inch))
        
        # Add a table
        data = [
            ['Name', 'Role', 'Department'],
            ['John Doe', 'Engineer', 'R&D'],
            ['Jane Smith', 'Manager', 'Sales'],
            ['Bob Johnson', 'Analyst', 'Finance']
        ]
        
        table = Table(data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(table)
        elements.append(Spacer(1, 0.3*inch))
        
        elements.append(Paragraph(
            "This PDF helps verify that the extraction system can properly parse PDF documents "
            "and extract both text and structured data.",
            normal_style
        ))
        
        # Build PDF
        doc.build(elements)
        print("✅ Created sample PDF")
        
    except ImportError:
        print("⚠️  reportlab not installed. Skipping PDF creation.")
        print("   Install with: pip install reportlab")


def main():
    """Generate all test files"""
    print("📁 Creating test files directory...")
    test_dir = create_test_files_directory()
    
    print("\n🔨 Generating test files...")
    create_sample_txt(test_dir)
    print("✅ Created sample TXT")
    
    create_sample_csv(test_dir)
    print("✅ Created sample CSV")
    
    create_sample_html(test_dir)
    print("✅ Created sample HTML")
    
    create_sample_image_with_text(test_dir)
    create_sample_docx(test_dir)
    create_sample_xlsx(test_dir)
    create_sample_pdf(test_dir)
    
    print(f"\n✨ Test files created in: {test_dir.absolute()}")
    print("\n📋 Generated files:")
    for file in sorted(test_dir.glob("*")):
        print(f"   - {file.name}")


if __name__ == "__main__":
    main()